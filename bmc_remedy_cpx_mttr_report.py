#!/usr/bin/env python3
"""
BMC Remedy (Helix ITSM) - CPX Incident MTTR Report via REST API

Pulls every incident from the HPD:Help Desk form whose Incident Number
starts with "CPX" and whose Submit Date falls in July, August or September,
then writes an Excel (.xlsx) report with:

    Incidents sheet  -> Incident Number, Summary, Status, Priority,
                        Reported Date, Submit Date, Last Resolved Date,
                        Time to Detect, Time to Resolve
    Summary sheet    -> per-month and overall incident count, resolved
                        count, Average MTTR and Average Time to Detect

Definitions:

    Time to Detect  (TTD)  = Submit Date        - Reported Date
    Time to Resolve (TTR)  = Last Resolved Date - Submit Date
    MTTR                   = average TTR of the resolved incidents
                             (open incidents have no Last Resolved Date
                             and are left out of the average)

Authentication: AR-JWT token from POST /api/jwt/login, sent as
"Authorization: AR-JWT <token>"; the token is released on exit.

Docs: https://docs.bmc.com -> Remedy AR System REST API -> Entry endpoints

Requirements:
    pip install requests openpyxl

Environment variables:

    BMC_HOST          e.g. "https://remedy.example.com:8443" (no trailing slash)
    BMC_USERNAME      Remedy user with read access to HPD:Help Desk
    BMC_PASSWORD      that user's password (prompted for if not set)
    BMC_PREFIX        optional, defaults to "CPX"
    BMC_PREFIX_FIELD  optional, field the prefix is matched on,
                      defaults to "Incident Number"
    BMC_DETECT_FIELD  optional, start field for Time to Detect,
                      defaults to "Reported Date"
    BMC_TIMEZONE      optional IANA zone for month boundaries and the dates
                      in the report, e.g. "Asia/Kolkata"; defaults to the
                      machine's local zone
    BMC_VERIFY_TLS    optional, "true"/"false", defaults to "true"

Run:

    python3 bmc_remedy_cpx_mttr_report.py
    python3 bmc_remedy_cpx_mttr_report.py --year 2026 --months 7 8 9
    python3 bmc_remedy_cpx_mttr_report.py --output cpx_q3.xlsx

    Without --year the current year is used.
"""

import argparse
import calendar
import getpass
import os
import sys
from datetime import datetime
from zoneinfo import ZoneInfo

import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

FORM = "HPD:Help Desk"
PAGE_SIZE = 500

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
TOTAL_FONT = Font(bold=True)
DATE_FORMAT = "yyyy-mm-dd hh:mm:ss"


# --------------------------------------------------------------------------- #
# Remedy REST API
# --------------------------------------------------------------------------- #

def login(host, username, password, verify):
    resp = requests.post(
        f"{host}/api/jwt/login",
        data={"username": username, "password": password},
        headers={"Content-Type": "application/x-www-form-urlencoded"},
        verify=verify,
        timeout=60,
    )
    resp.raise_for_status()
    return resp.text.strip()


def logout(host, token, verify):
    try:
        requests.post(
            f"{host}/api/jwt/logout",
            headers={"Authorization": f"AR-JWT {token}"},
            verify=verify,
            timeout=30,
        )
    except requests.RequestException:
        pass


def fetch_incidents(host, token, qualification, fields, verify):
    """Return every entry matching the qualification, following pagination."""
    url = f"{host}/api/arsys/v1/entry/{requests.utils.quote(FORM)}"
    headers = {"Authorization": f"AR-JWT {token}"}
    entries = []
    offset = 0
    while True:
        params = {
            "q": qualification,
            "fields": f"values({','.join(fields)})",
            "sort": "Submit Date.asc",
            "limit": PAGE_SIZE,
            "offset": offset,
        }
        resp = requests.get(url, headers=headers, params=params,
                            verify=verify, timeout=120)
        resp.raise_for_status()
        page = resp.json().get("entries", [])
        entries.extend(e.get("values", {}) for e in page)
        print(f"  fetched {len(entries)} incident(s)...")
        if len(page) < PAGE_SIZE:
            return entries
        offset += PAGE_SIZE


def build_qualification(prefix_field, prefix, start, end):
    # Date fields are compared as epoch seconds, which avoids any dependence
    # on the server's date-format locale.
    return (
        f"'{prefix_field}' LIKE \"{prefix}%\" "
        f"AND 'Submit Date' >= {int(start.timestamp())} "
        f"AND 'Submit Date' < {int(end.timestamp())}"
    )


# --------------------------------------------------------------------------- #
# Date / duration helpers
# --------------------------------------------------------------------------- #

def parse_remedy_date(value, tz):
    """Remedy returns e.g. "2026-07-03T14:22:10.000+0000"; result is naive local."""
    if not value:
        return None
    if isinstance(value, (int, float)):
        dt = datetime.fromtimestamp(value, tz)
    else:
        text = str(value).strip()
        for fmt in ("%Y-%m-%dT%H:%M:%S.%f%z", "%Y-%m-%dT%H:%M:%S%z"):
            try:
                dt = datetime.strptime(text, fmt)
                break
            except ValueError:
                continue
        else:
            dt = datetime.fromisoformat(text.replace("Z", "+00:00"))
            if dt.tzinfo is None:
                dt = dt.replace(tzinfo=tz)
    # Excel cannot store time zones, so write wall-clock time in the report zone.
    return dt.astimezone(tz).replace(tzinfo=None)


def hours_between(start, end):
    if start is None or end is None:
        return None
    return round((end - start).total_seconds() / 3600, 2)


def fmt_duration(hours):
    """12.5 -> "0d 12:30"; None -> "" (Excel-friendly text)."""
    if hours is None:
        return ""
    sign = "-" if hours < 0 else ""
    minutes = int(round(abs(hours) * 60))
    days, minutes = divmod(minutes, 24 * 60)
    hh, mm = divmod(minutes, 60)
    return f"{sign}{days}d {hh:02d}:{mm:02d}"


def average(values):
    values = [v for v in values if v is not None]
    return round(sum(values) / len(values), 2) if values else None


# --------------------------------------------------------------------------- #
# Report
# --------------------------------------------------------------------------- #

def build_rows(entries, detect_field, tz):
    rows = []
    for e in entries:
        reported = parse_remedy_date(e.get(detect_field), tz)
        submitted = parse_remedy_date(e.get("Submit Date"), tz)
        resolved = parse_remedy_date(e.get("Last Resolved Date"), tz)
        rows.append({
            "incident": e.get("Incident Number"),
            "summary": e.get("Description"),
            "status": e.get("Status"),
            "priority": e.get("Priority"),
            "reported": reported,
            "submitted": submitted,
            "resolved": resolved,
            "ttd": hours_between(reported, submitted),
            "ttr": hours_between(submitted, resolved),
            "month": submitted.strftime("%B %Y") if submitted else "",
        })
    rows.sort(key=lambda r: (r["submitted"] or datetime.min, r["incident"] or ""))
    return rows


def style_header(ws, headers):
    ws.append(headers)
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center",
                                   wrap_text=True)
    ws.freeze_panes = "A2"


def autosize(ws, max_width=60):
    for col in ws.columns:
        width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = \
            min(max(width + 2, 12), max_width)


def write_report(rows, months, year, detect_field, prefix, path):
    wb = Workbook()

    # ---- Incidents sheet ------------------------------------------------- #
    ws = wb.active
    ws.title = "Incidents"
    style_header(ws, [
        "Incident Number", "Summary", "Status", "Priority", "Month",
        detect_field, "Submit Date", "Last Resolved Date",
        "Time to Detect (hrs)", "Time to Detect (d hh:mm)",
        "Time to Resolve (hrs)", "Time to Resolve (d hh:mm)",
    ])
    for r in rows:
        ws.append([
            r["incident"], r["summary"], r["status"], r["priority"], r["month"],
            r["reported"], r["submitted"], r["resolved"],
            r["ttd"], fmt_duration(r["ttd"]),
            r["ttr"], fmt_duration(r["ttr"]),
        ])
    for row in ws.iter_rows(min_row=2, min_col=6, max_col=8):
        for cell in row:
            cell.number_format = DATE_FORMAT

    if rows:
        last = ws.max_row
        ws.append([])
        avg_row = ws.max_row + 1
        # Live formulas so the averages stay correct if rows are edited.
        ws.append([
            "Average", None, None, None, None, None, None, None,
            f'=IFERROR(ROUND(AVERAGE(I2:I{last}),2),"")', None,
            f'=IFERROR(ROUND(AVERAGE(K2:K{last}),2),"")', None,
        ])
        for cell in ws[avg_row]:
            cell.font = TOTAL_FONT
        ws.auto_filter.ref = f"A1:L{last}"
    autosize(ws)

    # ---- Summary sheet --------------------------------------------------- #
    ss = wb.create_sheet("Summary")
    style_header(ss, [
        "Period", "Total Incidents", "Resolved", "Open / Unresolved",
        "Average MTTR (hrs)", "Average MTTR (d hh:mm)",
        "Average Time to Detect (hrs)", "Average Time to Detect (d hh:mm)",
    ])

    def summary_line(label, subset):
        mttr = average(r["ttr"] for r in subset)
        mttd = average(r["ttd"] for r in subset)
        resolved = sum(1 for r in subset if r["resolved"])
        ss.append([label, len(subset), resolved, len(subset) - resolved,
                   mttr, fmt_duration(mttr), mttd, fmt_duration(mttd)])
        return mttr

    for m in months:
        label = f"{calendar.month_name[m]} {year}"
        summary_line(label, [r for r in rows if r["month"] == label])
    overall_label = (f"Overall ({calendar.month_abbr[months[0]]}-"
                     f"{calendar.month_abbr[months[-1]]} {year})")
    overall_mttr = summary_line(overall_label, rows)
    for cell in ss[ss.max_row]:
        cell.font = TOTAL_FONT

    ss.append([])
    ss.append([f"Incidents whose Incident Number starts with \"{prefix}\", "
               f"filtered on Submit Date."])
    ss.append([f"Time to Detect = Submit Date - {detect_field}; "
               f"Time to Resolve = Last Resolved Date - Submit Date."])
    ss.append(["MTTR averages resolved incidents only; open incidents "
               "are counted but excluded from the average."])
    autosize(ss, max_width=40)

    wb.move_sheet("Summary", offset=-1)
    wb.active = 0
    wb.save(path)
    return overall_mttr


# --------------------------------------------------------------------------- #
# Main
# --------------------------------------------------------------------------- #

def main():
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("--year", type=int, default=datetime.now().year,
                        help="year to report on (default: current year)")
    parser.add_argument("--months", type=int, nargs="+", default=[7, 8, 9],
                        help="month numbers to include (default: 7 8 9)")
    parser.add_argument("--output",
                        help="output .xlsx path (default: timestamped name)")
    args = parser.parse_args()

    host = os.environ.get("BMC_HOST", "").rstrip("/")
    username = os.environ.get("BMC_USERNAME")
    if not host or not username:
        sys.exit("Set BMC_HOST and BMC_USERNAME (see the header of this script).")
    password = os.environ.get("BMC_PASSWORD") or getpass.getpass("Remedy password: ")
    prefix = os.environ.get("BMC_PREFIX", "CPX")
    prefix_field = os.environ.get("BMC_PREFIX_FIELD", "Incident Number")
    detect_field = os.environ.get("BMC_DETECT_FIELD", "Reported Date")
    verify = os.environ.get("BMC_VERIFY_TLS", "true").lower() != "false"
    tz_name = os.environ.get("BMC_TIMEZONE")
    tz = ZoneInfo(tz_name) if tz_name else datetime.now().astimezone().tzinfo

    months = sorted(set(args.months))
    if any(m < 1 or m > 12 for m in months):
        sys.exit("--months must be between 1 and 12.")
    output = args.output or (f"{prefix}_incident_mttr_report_"
                             f"{datetime.now():%Y%m%d_%H%M%S}.xlsx")

    fields = ["Incident Number", "Description", "Status", "Priority",
              "Submit Date", "Last Resolved Date", detect_field]

    token = login(host, username, password, verify)
    try:
        entries = []
        for m in months:
            start = datetime(args.year, m, 1, tzinfo=tz)
            end = (datetime(args.year + 1, 1, 1, tzinfo=tz) if m == 12
                   else datetime(args.year, m + 1, 1, tzinfo=tz))
            print(f"Querying {calendar.month_name[m]} {args.year}...")
            entries.extend(fetch_incidents(
                host, token,
                build_qualification(prefix_field, prefix, start, end),
                fields, verify,
            ))
    finally:
        logout(host, token, verify)

    rows = build_rows(entries, detect_field, tz)
    overall_mttr = write_report(rows, months, args.year, detect_field,
                                prefix, output)

    print(f"\n{len(rows)} {prefix} incident(s) found.")
    print(f"Average MTTR: {fmt_duration(overall_mttr) or 'n/a'}"
          + (f" ({overall_mttr} hrs)" if overall_mttr is not None else ""))
    print(f"Report written to {output}")


if __name__ == "__main__":
    main()
